"""Fast, title-only regression comparison; references never supply production values."""

from __future__ import annotations

import argparse
from collections import Counter
import importlib.util
import json
from pathlib import Path
import time

from validate import check, risk_categories

ROOT = Path(__file__).resolve().parent


def read_reference(path):
    # Physical JSONL lines preserve U+2028 inside an original title.
    with Path(path).open(encoding='utf-8') as file:
        records = [json.loads(line) for line in file if line.strip()]
    if not records or len({item['row'] for item in records}) != len(records):
        raise ValueError('Reference must contain nonempty, unique original row IDs')
    return records


def load_extractor(path):
    spec = importlib.util.spec_from_file_location('experiment_extract', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def category(result, reference, failures=()):
    if failures:
        return 'invalid_script'
    if reference['status'] in ('uncertain', 'conflict'):
        return 'unresolved_reference'
    if (result['status'], result['names']) == (reference['status'], reference['names']):
        return 'match'
    if result['status'] == 'uncertain':
        return 'script_unresolved_confirmed_reference'
    if len(result['names']) != len(reference['names']):
        return 'count_or_split_difference'
    return 'name_or_status_difference'


def evaluate(records, extractor):
    started = time.perf_counter()
    cache, errors, production = {}, [], []
    categories, statuses, error_rules, risks = Counter(), Counter(), Counter(), Counter()
    failure_rows = 0
    for reference in records:
        title = reference['title']
        if title not in cache:
            result = extractor.extract(title)
            result = dict(result, length=sum(map(extractor.name_length, result['names'])),
                          extracted_name='；'.join(result['names']))
            cache[title] = result, check(title, result)
        result, checks = cache[title]
        state = category(result, reference, checks['failures'])
        categories[state] += 1
        statuses[result['status']] += 1
        failure_rows += bool(checks['failures'])
        if state != 'match':
            tags = risk_categories(title, result, checks)
            error_rules[result['rule']] += 1
            risks.update(tags)
            errors.append(dict(row=reference['row'], input_id=reference.get('input_id'),
                               title=title, category=state, result=result,
                               reference={key: reference[key] for key in
                                          ('status', 'names', 'source_quality') if key in reference},
                               checks=checks, risk_categories=tags))
        # Abstention remains blank/blank. AI/reference answers never fill cells.
        production.append(dict(row=reference['row'],
                               name=None if result['status'] == 'uncertain' else result['extracted_name'],
                               length=None if result['status'] == 'uncertain' else result['length']))
    rows = len(records)
    confirmed = rows - categories['unresolved_reference']
    determined = rows - statuses['uncertain']
    determined_matches = sum(category(result, reference, checks['failures']) == 'match'
                             for reference in records
                             for result, checks in [cache[reference['title']]]
                             if result['status'] != 'uncertain')
    summary = dict(script_version=extractor.VERSION, source_rows=rows, unique_titles=len(cache),
                   categories_rows=dict(categories), status_rows=dict(statuses),
                   confirmed_matching_rows=categories['match'],
                   confirmed_difference_rows=confirmed - categories['match'],
                   unresolved_reference_rows=categories['unresolved_reference'],
                   reference_agreement=categories['match'] / rows,
                   confirmed_reference_agreement=categories['match'] / confirmed if confirmed else None,
                   script_determinate_rows=determined,
                   script_determinate_reference_agreement=determined_matches / determined if determined else None,
                   source_information_conflict_rows=sum(
                       item.get('source_quality') == 'title_body_information_conflict' for item in records),
                   contract_failure_rows=failure_rows,
                   contract_failure_titles=sum(bool(checks['failures']) for _, checks in cache.values()),
                   errors_by_rule=dict(error_rules.most_common()), errors_by_risk=dict(risks.most_common()),
                   unseen_data_accuracy=None, true_accuracy=None, acceptance_passed=False,
                   seconds=round(time.perf_counter() - started, 3))
    return summary, errors, production


def check_source(source, records):
    from openpyxl import load_workbook
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = book['Sheet1']
        if sheet.cell(1, 1).value != '标题':
            raise ValueError('Expected the title field in column A of Sheet1')
        titles = {row: title for row, (title,) in enumerate(sheet.iter_rows(
            min_row=2, max_row=max(item['row'] for item in records),
            min_col=1, max_col=1, values_only=True), 2)}
        for item in records:
            if titles.get(item['row']) != item['title']:
                raise ValueError(f'Original title/order differs at Excel row {item["row"]}')
    finally:
        book.close()
    return {'original_titles_verified': len(records), 'title_column': 'A', 'sheet': 'Sheet1'}


def evaluate_unseen(path, regression, extractor):
    unseen = read_reference(path)
    seen_titles = {item['title'] for item in regression}
    if any(item['title'] in seen_titles for item in unseen):
        raise ValueError('Unseen references overlap the fixed regression titles')
    for item in unseen:
        if item['status'] not in ('names', 'no_name') or check(item['title'], item)['failures']:
            raise ValueError('Unseen accuracy requires resolved, title-grounded independent labels')
    return evaluate(unseen, extractor)


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def write_errors(path, errors):
    with path.open('w', encoding='utf-8') as file:
        for item in errors:
            file.write(json.dumps(item, ensure_ascii=False, separators=(',', ':')) + '\n')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--reference', type=Path, default=ROOT / 'reference_v3.jsonl')
    parser.add_argument('--extractor', type=Path, default=ROOT / 'extract.py')
    parser.add_argument('--source', type=Path, default=ROOT / 'source.xlsx')
    parser.add_argument('--output', type=Path)
    parser.add_argument('--check-source', action='store_true')
    parser.add_argument('--export', type=Path, help='Append script-only N/O columns to an original workbook copy')
    parser.add_argument('--unseen-reference', type=Path,
                        help='Independent, never-used title labels in the same JSONL format')
    args = parser.parse_args()
    records = read_reference(args.reference)
    extractor = load_extractor(args.extractor.resolve())
    summary, errors, production = evaluate(records, extractor)
    summary['reference_version'] = 'v3' if args.reference.name == 'reference_v3.jsonl' else args.reference.stem
    summary['reference_file'] = args.reference.name
    summary['extractor_file'] = str(args.extractor.resolve())
    if summary['reference_version'] == 'v3':
        summary['v3_reference_agreement'] = summary['reference_agreement']
    if args.check_source or args.export:
        summary['source_check'] = check_source(args.source, records)
    unseen_result = None
    if args.unseen_reference:
        unseen_summary, unseen_errors, _ = evaluate_unseen(args.unseen_reference, records, extractor)
        summary['unseen_data_accuracy'] = unseen_summary['reference_agreement']
        summary['unseen_rows'] = unseen_summary['source_rows']
        unseen_result = unseen_summary, unseen_errors
    if args.export:
        if summary['contract_failure_rows']:
            raise ValueError('Script output contract failed; export stopped')
        from excel_output import export_workbook
        summary['excel_export_check'] = export_workbook(args.source, args.export, production)
    output = args.output or ROOT / 'results' / extractor.VERSION
    output.mkdir(parents=True, exist_ok=True)
    write_json(output / 'summary.json', summary)
    write_errors(output / 'errors.jsonl', errors)
    if unseen_result:
        write_json(output / 'unseen-summary.json', unseen_result[0])
        write_errors(output / 'unseen-errors.jsonl', unseen_result[1])
    agreement = f'{summary["reference_agreement"]:.4%}'
    unseen_value = summary['unseen_data_accuracy']
    unseen_text = '未提供独立新数据，尚未测量' if unseen_value is None else f'{unseen_value:.4%}'
    report = [f'# {extractor.VERSION} 评测', '',
              f'固定参考：{summary["reference_version"]}；{summary["source_rows"]:,} 原始行，'
              f'{summary["unique_titles"]:,} 个标题。', '',
              f'参考一致率 **{agreement}**，一致 {summary["confirmed_matching_rows"]:,} 行；'
              f'未见数据准确率：{unseen_text}。', '',
              '固定参考已参与历史规则与词库开发，只用于回归；真实准确率和最终验收未确认。', '',
              '| 分类 | 原始行数 |', '| --- | ---: |']
    report.extend(f'| {state} | {count:,} |' for state, count in summary['categories_rows'].items())
    report.extend(['', '错误分析见 `errors.jsonl`，含原标题、参考、脚本结果、规则与风险。', '',
                   '| 差异规则 | 原始行数 |', '| --- | ---: |'])
    report.extend(f'| {rule} | {count:,} |' for rule, count in list(summary['errors_by_rule'].items())[:15])
    (output / 'report.md').write_text('\n'.join(report) + '\n', encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
