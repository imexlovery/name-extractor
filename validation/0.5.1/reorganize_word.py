"""Put current r56 rules first, preserving the former document as history.

Run with the bundled document runtime. Read the retained pre-edit copy on rerun.
"""
from copy import deepcopy
import json
from pathlib import Path
import shutil

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parents[2]
PATH = ROOT / '任职资格批复姓名提取与自动验证方案.docx'
BACKUP = ROOT / '.local/word-history/20261012-before-template-semantic.docx'
OUT = Path(__file__).resolve().parent / 'template-semantic'


def main():
    summary = json.loads((OUT / 'summary.json').read_text())
    sampling = json.loads((OUT / 'selection-summary.json').read_text())
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(PATH, BACKUP)
    original = Document(BACKUP)
    doc = Document(BACKUP)
    old_blocks = [deepcopy(e) for e in original.element.body if e.tag != qn('w:sectPr')]
    for e in list(doc.element.body):
        if e.tag != qn('w:sectPr'):
            doc.element.body.remove(e)

    def paragraph(text, style=None):
        return doc.add_paragraph(text, style)

    def heading(text):
        return paragraph(text, 'Heading 1')

    paragraph('任职资格批复姓名提取与自动验证方案', 'Title')
    paragraph('2026年10月12日  0.5.1  r56  现行规则与独立语义验证')
    paragraph('本方案指导Codex仅依据原标题提取任职资格批复中的明示姓名。前半部分集中说明当前有效规则、运行与验证边界；后半部分完整保留原方案的阶段记录，便于追踪修订原因。当前生产规则继续使用r56；本次新增分层盲标用于主动寻找未知错误，结论见现行验证状态。')
    heading('数据范围与输出口径')
    for i in (4, 5, 6, 8, 9, 10):
        paragraph(original.paragraphs[i].text)
    example_table = original.tables[0]._element
    doc.element.body.insert(len(doc.element.body) - 1, deepcopy(example_table))
    heading('标题结构与完整主体判断')
    paragraphs = [
        '首先识别被批复人的完整任职资格或岗位语义，联合判断监管前缀、许可动作、机构、姓名、完整岗位及后缀。拟任资格和明确董事资格可以提供边界；标题中引用管理规定的任职资格不能抢占书名号外的具名许可事项。规范性通知或行政文书的类型不自动证明没有明示人员。裸核准对象没有任职、岗位或称谓语义时保留未决，不凭短字符片段认定姓名。',
        '机构锚点仅提供机构边界，不是姓名或行号答案表。当前加载institution_anchors_r41.json的12271个机构片段及82个基础锚点，合计12353；每项原标题来源保留。完整机构及母机构与分支后缀联合约束姓名位置，弱地区词不得抢占完整机构。机构前主体与机构岗位后主体构成不同的完整解释时未决。未知机构边界不靠常见姓氏或固定姓名长度猜切。',
        '岗位从最长完整表达开始识别，避免从财务会计部、财务总监、风险官、高级管理人员等表达中间切出残片。姓名的任职资格中的结构助词的应剥离；核字缺少准字或拟字处于姓名与任职结构交界时，若有许可动作残片与完整姓名两种解释则未决。只有逐字相同的重复标题或尾句才按重复内容处理，不因此丢弃不同人员。',
        '括号外的顿号、分号等明示分项先形成子句，每项分别判断为姓名主体、依附前人岗位或机构表头，再合并全部明示姓名并保留原标题顺序。完整岗位后的和、及、与在r56中引出后续子句，必须一并检查；姓名、机构或部门内部的连接字不独自触发拆分。只有岗位的后项可以依附同一人，兼任不自动变成多人。',
        '完整岗位后仍有未消费的括号外内容时，保留候选并未决，不因首个岗位已识别而忽略后续主体。没有可靠分隔的连写姓名与岗位结构不靠姓氏猜切。完整营业部总经理助理、理事会秘书和等字省略按完整岗位及既有语义处理。无姓名判断须检查完整标题及全部剩余岗位、履职括注和主体，机构语境本身不能证明无姓名。',
        '机构开业并核准人员时，先消费明确开业事项，再递归处理整个许可子句。明确机构表头后的冒号或双破折号可以引出名单；括号内只有明示岗位名单及冒号的内容按名单解析，普通别名、履职备注和数字说明不能自动当作另一主体。书名号和括号保留原位置，避免机构括注干扰姓名边界。',
        '同人的中英文表达与括号别名保留原文，括号内逗号和连字符不拆分。括号外外文逗号不足以证明同人或多人时，保留完整候选并未决。中点长姓名保留整段；无明确边界的长连续汉字不猜拆。候选边缘只去除不可见格式字符，姓名内部及未定义字符的歧义继续披露。',
    ]
    for text in paragraphs:
        paragraph(text)
    heading('姓名人数与明确弃权')
    for i in (97, 98, 99, 100):
        paragraph(original.paragraphs[i].text)
    heading('独立验证与真实错误处置')
    paragraph('A和B在独立上下文中逐标题理解，只接收原标题、记录标识及共同业务口径，不读取脚本答案、旧参考或另一通道结果。各自保留原始首答、标题依据与人数证据；更正须另存before、after及理由，不能覆盖首答或以多数票关闭语义分歧。相同模型的独立上下文仍可能共享偏差，不能作为人工真值。')
    for i in (20, 21, 32, 34):
        paragraph(original.paragraphs[i].text)
    paragraph('内部台账保留原行映射、选样种子、规则版本、首答和更正依据，重启复用已有参考。错误修订先归类共同原因并同步Word与代码，再全量比较姓名、展示值、字数及状态。故障注入、合成契约、重复运行和导出不计真实数据错误或真实规则修改轮次。')
    paragraph('当前新增验证从缺参考真实标题中按模板稀有度、姓名槽位长度、语法边界字和低频字符选样。稀有模板尽量全覆盖；高频模板兼顾固定种子随机样本与边界样本，总量由覆盖条件去重产生。槽位检查只用于选样和机械检查，不生成AI姓名答案，不证明姓名或人数语义。只有真实标题中的错误依据确认后才追踪共同根因、修订通用规则并全量回归；没有错误证据时保留r56。')
    heading('现行验证状态')
    paragraph('r56全量为46760个不同标题、52146条原行；姓名45944标题51254行、无姓名694标题761行、未决122标题131行。相对r55姓名及状态变化0，原全量输出合同失败0；本次不修改生产规则、机构资源或弃权名单，不重新导出输出值相同的Excel，继续使用0.5.0待复核版。')
    states = summary['first_answer_comparison_states']
    paragraph(f"新增验证的冻结缺参考池为{sampling['missing_reference_titles']}标题，归为{sampling['templates']}种直接结构模板；二字槽位7021、三字15789，没有标点或外文槽位。稀有阈值为每模板不超过12标题，共{sampling['rare_templates']}种、{sampling['rare_template_titles']}标题全部纳入；其余模板按频率和边界特征选样，去重后{summary['selected_titles']}标题、{summary['selected_mapped_rows']}原行，两路共{2 * summary['selected_titles']}份首答。")
    paragraph(f"两路首答一致{summary['first_answers_agree']}标题，其中{states.get('match', 0)}标题与冻结r56输出匹配，确定差异{states.get('determinate_difference', 0)}、参考分歧{states.get('disputed_reference', 0)}。标题明示人数证据A为{summary['explicit_person_count_evidence']['a']}、B为{summary['explicit_person_count_evidence']['b']}；完整单槽不是内部单人数证明。新增参考后的缺参考标题为{summary['full_comparison_states_with_new_labels'].get('missing_reference', 0)}；旧单渠道和语义分歧继续保留。目的性抽样不能估计全量真实准确率，AI一致也不能代替人工检查。")
    paragraph('现有r56基线与结构审查快照保留原统计，新模板统计、选样、独立首答和覆盖变化单独保存于validation/0.5.1/template-semantic及labels。真实准确率未评估，最终人工检查未执行，不能宣称99.8%或100%达标、全量语义通过或最终验收通过。')
    heading('运行与证据入口')
    paragraph('生产使用uv和项目本地.venv。全量运行命令为uv run python -B scripts/full.py run，参考比较为uv run python -B scripts/full.py finish；两者会更新默认工作目录中的阶段结果，需保留既有冻结材料。新增分层选样复现为uv run python -B validation/0.5.1/template_semantic.py；首答完成后用同一命令加evaluate生成独立比较报告。AI答案不参与生产提取或Excel回填。')
    paragraph('项目要求.md是最高参考，维护文档仅记录状态和方案依据。现行代码位于scripts，r56源码及实际资源快照位于validation/0.5.1；新增验证报告和模板统计与历史回归分开保存。没有真实错误证据不增加规则版本，已揭示的参考以后属于工程回归，不再称为新盲测。')
    appendix = heading('历史迭代记录')
    appendix.paragraph_format.page_break_before = True
    paragraph('以下保留原方案第一节至第十七节及逐版修订表。历史文字中的当前、下一步、统计和章节引用均以各阶段当时的口径理解；其中与现行规则不同的旧策略不得直接用于生产。阶段验证、合成契约和机械检查均不等同于最终人工验收。')
    for e in old_blocks[3:]:
        if e.tag == qn('w:tbl') and ''.join(e.itertext()) == ''.join(example_table.itertext()):
            continue
        style = e.find('./' + qn('w:pPr') + '/' + qn('w:pStyle'))
        if style is not None and style.get(qn('w:val')) == 'Heading1':
            style.set(qn('w:val'), 'Heading2')
        doc.element.body.insert(len(doc.element.body) - 1, e)
    for name in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.save(PATH)
    result = Document(PATH)
    old_texts = [p.text for p in original.paragraphs[3:] if p.text]
    result_texts = [p.text for p in result.paragraphs]
    assert all(text in result_texts for text in old_texts), 'Historical paragraph lost'
    assert len(result.tables) == len(original.tables), 'Historical table lost or duplicated'
    (OUT / 'word-reorganization.json').write_text(json.dumps(dict(
        original_paragraphs=len(original.paragraphs), output_paragraphs=len(result.paragraphs),
        historical_paragraphs_preserved=len(old_texts), tables_preserved=len(result.tables),
        retained_pre_edit_copy=str(BACKUP), current_rules_before_history=True,
        rendering_verified=False,
    ), ensure_ascii=False, indent=2) + '\n')
    print(PATH)


if __name__ == '__main__':
    main()
