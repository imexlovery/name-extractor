"""Preserve human feedback and freeze title-only human review lists before edits."""
import csv
import json
from pathlib import Path
import random
import re
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts'))
from pilot import read_jsonl, write_json, write_jsonl

WORK = Path(__file__).resolve().parent


def review_csv(path, items):
    with path.open('w', encoding='utf-8-sig', newline='') as file:
        writer = csv.writer(file)
        writer.writerow(['序号', '完整标题', '人工姓名或无姓名判断', '判断依据及是否查看正文', '原Excel行号'])
        for index, item in enumerate(items, 1):
            writer.writerow([index, item['title'], '', '', '；'.join(map(str, item['rows']))])


def main():
    baseline = read_jsonl(ROOT / 'validation/0.5.1/script-r56.jsonl')
    titles = {v['title']: v for v in baseline}
    source = ROOT / 'outputs/01a13013-4021-7ca1-8cf8-969859ee2d2d/r56-122个未决标题-人工盲审.xlsx'
    workbook = load_workbook(source, read_only=True, data_only=True)
    confirmed = {v['row']: v for v in json.loads((WORK / 'user-confirmed-corrections.json').read_text())}
    feedback = []
    for index, title, first, note, correction, rows in workbook.active.iter_rows(min_row=2, values_only=True):
        item = titles[title]
        mapped = [int(v) for v in re.split('[；;,，]', str(rows))]
        assert mapped == item['rows'] and item['result']['status'] == 'uncertain'
        updated = correction or first
        status = 'no_name' if updated == '无' else 'uncertain' if updated in ('不知道', '标题不完整', '不完整') else 'names'
        names = re.split('[；;]', updated) if status == 'names' else []
        # The prior conversation explicitly verified six corrections and five foreign
        # forms against content. This records that source, never a runtime answer.
        content_checked = bool(correction) or index in (73, 84, 89, 95, 101) or bool(note and re.search('正文|内容|查了|核查', note))
        family = ('foreign_comma' if '外文括号外' in item['result']['note'] else
                  'connector' if '连接词关系' in item['result']['note'] else
                  'long_han' if '较长连续汉字' in item['result']['note'] else item['result']['rule'])
        record = dict(index=index, row=item['row'], rows=mapped, title=title,
                      original_c=first, original_d=note, original_e=correction,
                      updated_judgment=dict(status=status, names=names),
                      evidence_source='正文辅助确认' if content_checked else '来源尚不明确',
                      outstanding_question=bool(note and re.search('待确认|待确定|无法确定|不确定|不知道|也许', note)),
                      family=family, baseline=item['result'], human_truth_independent=False)
        if index in (73, 84, 89, 95, 101):
            record['conversation_correction'] = dict(
                names=item['result']['names'], status='names', source='目标聊天及0.6.0计划第2.3/8节的正文核验',
                applies_to_production=False)
        if item['row'] in confirmed:
            record['user_confirmed_correction'] = confirmed[item['row']]
        feedback.append(record)
    workbook.close()
    assert len(feedback) == 122 and sum(len(v['rows']) for v in feedback) == 131
    write_jsonl(WORK / 'human-review.jsonl', feedback)
    # Reruns preserve the original, unexposed sample.
    if not (WORK / 'random-human-frozen.jsonl').exists():
        population = [v for v in baseline if v['result']['status'] == 'names']
        sample = random.Random(202610120600).sample(population, 200)
        sample = [dict(row=v['row'], rows=v['rows'], title=v['title']) for v in sample]
        write_jsonl(WORK / 'random-human-frozen.jsonl', sample)
        write_json(WORK / 'random-human-manifest.json', dict(seed=202610120600, population=len(population),
                   sample_size=200, sampling_unit='r56已提取姓名去重标题', used_for_tuning=False,
                   human_review_complete=False))
        review_csv(WORK / '随机200标题-人工盲审.csv', sample)
    no_name = [dict(row=v['row'], rows=v['rows'], title=v['title']) for v in baseline if v['result']['status'] == 'no_name']
    review_csv(WORK / '原694无姓名标题-人工盲审.csv', no_name)
    write_json(WORK / 'human-review-scope.json', dict(
        initial_review_titles=122, initial_review_rows=131, corrections=6,
        post_change_confirmed_titles=len(confirmed), post_change_confirmed_rows=list(confirmed),
        content_assisted=True, all_titles_human_reviewed=False,
        no_name_review_titles=694, no_name_review_rows=761, no_name_review_complete=False,
        random_review_titles=200, random_review_complete=False))
    print('Saved 122 feedback records, 694 no-name review titles, frozen 200 random titles.')


if __name__ == '__main__':
    main()
