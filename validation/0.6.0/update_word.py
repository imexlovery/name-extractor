"""Update only effective-rule paragraphs; retain the historical appendix."""
from pathlib import Path
import json
import shutil
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
PATH = ROOT / '任职资格批复姓名提取与自动验证方案.docx'
BACKUP = ROOT / '.local/word-history/20261012-before-0.6.0.docx'

updates = {
    1: '2026年10月12日  0.6.0  r60  人工反馈与标题结构改进',
    2: '本方案指导Codex仅依据原标题提取任职资格批复中的明示姓名。前半部分说明当前有效规则和验证边界，后半部分保留历史迭代记录。0.6.0依据原122个未决标题的人工反馈，改进完整机构、银行类型与明确并列关系；生产使用r60，输出仍待人工复核。',
    6: '原邮件要求最终人工检查。用户已填写原122个未决标题的复核表，其中部分判断查看过正文。初答、更正和疑问分别保留，正文不能成为生产输入。原694个无姓名标题、冻结200个随机标题及规则变化后的人工复核尚未完成，不能将部分人工填写或AI比较视为最终验收。',
    12: '机构锚点仅提供机构边界，不是姓名或行号答案表。当前institution_anchors_r57.json包含12278个机构片段及82个基础锚点，合计12360；新增7个完整机构的标题来源另存。完整机构优先于村镇银行或农村商业银行等类型后缀。无完整机构竞争时，姓名候选加完整类型及全部消费的明确岗位可建立泛称类型边界；纯类型加岗位仍为无姓名。逯晓利村镇银行的人工疑问继续未决。固定数据中可用的类型规则不能保证对未来陌生银行品牌泛化，未知机构不靠姓名长度猜起点。',
    14: '括号外的顿号、分号先形成完整子句，逐项判断姓名主体、依附岗位或机构表头，再按原序合并。姓名槽位中的和、及、与可以引出并列主体，但每侧须通过既有完整片段检查；连接词只在首尾留下单字或含机构岗位残片时仍未决。带外文别名的整体不在此规则中拆分，保护坂内和隆（SAKAUCHIKAZUTAKA）等表达。完整岗位后的连接词仍消费全部子句；兼任不自动成为多人，明确一人可保留含连接字的整段姓名。',
    22: 'title_ambiguities_r57.json继承原7条弃权记录，并新增逯晓利村镇银行的人工疑问，共3个真实标题和5个历史合成题。它只保留无法消除的竞争解释，不提供姓名答案。完整标题命中才生效；新增精确人数仍由通用规则判断。王夏春子、殷熊毅辉的正文或人工判断另存，不能据此删除人数争议。',
    28: '报告分别披露规则状态、输出合同、脚本与已有AI参考比较及实际人工范围。当前人工记录涵盖原122个未决标题；变化后检查、原694个无姓名标题和冻结200个随机标题仍待人工完成。没有总体可靠真值时真实准确率记为未评估，不能宣称99.8%或100%验收通过。',
    30: '0.6.0先冻结200个r56已提取姓名去重标题，再试验机构与连接词。小范围单姓、复姓表只离线提出连写切分候选，整段始终保留为竞争解释；词表唯一切分不能证明人数，未进入生产。长连续汉字和外文逗号经过逐条归类，现有标题不足以支持可靠通用放行条件时保留候选；这些条目包括规则能力不足，不一概称为标题本身不可判断。',
    32: 'r60全量为46760个去重标题、52146条原行；姓名45991标题51303行、无姓名694标题761行、未决75标题82行。相对r56新增确定47标题49行，均与更正后的人工判断一致；用户已确认第25395行钱锦，人工初答误把理事首字带入姓名，更正独立保存。变化后检查仅完成该1条，其余仍待人工复核。原已确定输出变化0，合同失败0。r58的外文别名误拆已在r59修复；r60保护明示一人时的重复片段，真实输出未变。',
    33: '此前0.5.1模板验证新增285标题308行两路首答，285全部一致且匹配冻结r56。0.6.0复用此前所有阶段首答、显式更正及模板参考，不重复标注整个风险池，也不把旧参考计为新增。',
    34: 'r60与已有AI参考比较：匹配24121、脚本未决且有参考42、缺参考22525、单渠道35、参考分歧37，确定差异0。已有AI参考用于工程回归，不能当作人工总体真值；其中仍有SURIKOV，ARTEM等历史确定参考与当前弃权的差异。',
    35: '最终38项测试通过。源码、机构和弃权资源、逐轮结果、差异、人工原值及证据分类保存于validation/0.6.0。原表导出继续只追加提取姓名和姓名字数，另存0.6.0纯脚本待复核版；原13字段、52146行和格式通过独立导出检查后，仍需完成最终人工检查。',
    37: '使用uv与项目本地.venv。uv run python -B scripts/full.py run保存当轮源码及实际资源，同时比较上一轮和r56；finish复用已有参考且只生成script_only逐原行输出；export沿用原OOXML追加方式。离线人工比较使用uv run python -B validation/0.6.0/evaluate_stage.py。准备脚本冻结的随机样本不用于调参；未完成的人工检查入口为本阶段CSV。',
    38: '项目要求.md为最高参考。当前规则及机构资源位于scripts，r56原成果保留于validation/0.5.1；0.6.0报告、逐轮差异、反馈分类和剩余未决单独保存。仅支持当前有依据的局部修复，不能为发布版本或减少未决增加姓名回填。',
}


def main():
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(PATH, BACKUP)
    old = Document(BACKUP)
    doc = Document(BACKUP)
    old_history = [p.text for p in old.paragraphs[39:]]
    for index, text in updates.items():
        paragraph = doc.paragraphs[index]
        paragraph.clear()
        paragraph.add_run(text)
    doc.save(PATH)
    result = Document(PATH)
    assert [p.text for p in result.paragraphs[39:]] == old_history
    assert [t._element.xml for t in result.tables] == [t._element.xml for t in old.tables]
    (ROOT/'validation/0.6.0/word-update.json').write_text(json.dumps(dict(
        updated_effective_paragraphs=sorted(updates), historical_paragraphs_preserved=len(old_history),
        tables_preserved=len(result.tables), rendering_verified=False,
        pre_edit_copy=str(BACKUP)), ensure_ascii=False, indent=2)+'\n')
    print(PATH)


if __name__ == '__main__':
    main()
